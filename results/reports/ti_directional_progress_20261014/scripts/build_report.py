"""Build the Chinese progress report from the curated scientific JSON package.

Requires python-docx. It does not launch or modify any simulation.
Render the DOCX with the packaged document renderer before publication.
"""
from pathlib import Path
from datetime import datetime, timezone
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
FIG = ROOT / 'figures'
OUT = ROOT / 'report'
OUT.mkdir(exist_ok=True)

def read(name):
    return json.loads((DATA/name).read_text())

records = read('current_structured_results.json')['records']
flat = read('flat_validation_results.json')['records']
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.65), Inches(.65)
s.left_margin, s.right_margin = Inches(.72), Inches(.72)
s.header_distance, s.footer_distance = Inches(.25), Inches(.27)
styles = doc.styles
for name, size in [('Normal',10.5),('Title',24),('Subtitle',11),('Heading 1',17),('Heading 2',13),('Caption',9)]:
    st=styles[name]; st.font.name='Arial Unicode MS'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    fonts=st.element.get_or_add_rPr().rFonts
    for key in list(fonts.attrib):
        if 'Theme' in key or 'theme' in key:del fonts.attrib[key]
    fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    pp=st.element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.12
    st.font.italic=False
styles['Title'].paragraph_format.space_after=Pt(14)
styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Normal'].paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2','Caption']:
    styles[name].paragraph_format.keep_with_next = name != 'Caption'

# Page numbers help navigate this multi-section scientific report.
p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); r._r.addnext(field)
doc.core_properties.title='钛斜槽定向热发射计算阶段报告'
doc.core_properties.subject='同材料模型下的占空比与波长抽样结果及数值诊断'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.keywords='Meep; Ti; exploratory; directional thermal emission'
doc.core_properties.comments=''
doc.core_properties.created=datetime.now(timezone.utc)
doc.core_properties.modified=doc.core_properties.created

def p(text,style=None):
    return doc.add_paragraph(text,style)
def head(text):
    doc.add_heading(text,level=1)
def page(text):
    heading=doc.add_heading(text,level=1)
    heading.paragraph_format.page_break_before=True
def sub(text):
    doc.add_heading(text,level=2)
def table(headers, rows, widths=None, size=10):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top',75),('bottom',75),('left',95),('right',95)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
            tcPr.append(mar)
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8EDF0');tcPr.append(shd)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.12
                for run in para.runs:
                    run.font.size=Pt(size);run.bold=i==0
    p('')
    return t
def image(name,caption,width=7.0,max_height=None):
    path=FIG/name
    if not path.exists():raise FileNotFoundError(path)
    if max_height:
        with Image.open(path) as im:width=min(width,max_height*im.width/im.height)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(4)
    para.add_run().add_picture(str(path),width=Inches(width))
    p(caption,'Caption')
def val(fill,lam,theta):
    return next(x for x in records if x['fill']==fill and x['wavelength_um']==lam and x['theta_deg']==theta)
def pct(x):return f'{100*x:.3f}'
def pair(fill,lam):
    a,b=val(fill,lam,30),val(fill,lam,-30)
    return [a,b,100*(a['metrics']['A_flux']-b['metrics']['A_flux'])]

p('钛斜槽定向热发射计算阶段报告','Title')
p('截至 2026 年 10 月 14 日 13 时 45 分（北京时间）','Subtitle')
p('研究目标是建立特征尺寸对 8–13 μm 发射响应及方向性的影响规律，并在后续扩展至 3–20 μm。当前同一材料模型下已获得 10 个结构计算点：F=0.65 的 8、10.5、13 μm p 偏振 ±30°，以及 F=0.50、0.80 的 10.5 μm p 偏振 ±30°。')
p('三组占空比中，F=0.65 在 10.5 μm 的 +30° 吸收响应和正负角度差最大。通量法与体积积分法给出相同的三点非单调趋势，因此可以据此选择下一轮重点区域；目前仍属于探索结果，尚未完成带内、双偏振或全角域验证。')
sub('核心结果')
table(['开口占空比 F','+30° Aflux','−30° Aflux','正负角差'],[[f'{f:.2f}',pct(pair(f,10.5)[0]['metrics']['A_flux'])+'%',pct(pair(f,10.5)[1]['metrics']['A_flux'])+'%',f'{pair(f,10.5)[2]:.3f} 个百分点'] for f in [.5,.65,.8]],[1.35,1.65,1.65,2.35])
p('最新四项任务均已结束，并通过本轮探索性停止和一致性检查，耗时约 18 小时 35 分。完成状态为 EXPLORATORY_CONVERGED_NOT_QUALIFIED；“已完成”不等于正式数值资格认定。[D1]')
sub('结果的适用范围')
p('可支持：固定周期、轴长和倾角下的三点占空比筛选；固定 F=0.65 下的三个波长抽样比较；同一计算定义下的 ±30° 对比。')
p('尚不支持：F=0.65 是全局最优；8–13 μm 连续或带平均发射率；角度峰值与半高宽；非偏振性能；三维斜孔实物或高温样品性能。')
p('现阶段关键缺口是谱角覆盖和数值闭合，而不是缺少更大的盲目扫描。报告保留小负透射及通量与体积吸收的偏差，不将其裁剪或隐藏。')

page('1 计算模型与参数定义')
image('current/05_geometry.png','图 1  当前二维周期斜盲槽几何。F 为表面空气开口宽度与周期之比，不是钛的体积分数。示意使用保存的几何顶点重绘。[D4]',6.9)
table(['固定条件','当前设置'],[['几何与维度','二维周期斜盲槽；周期 50 μm；轴长 40 μm；倾角 30°；钛层总厚 60 μm'],['空间与时间离散','r20，即每 μm 20 格，网格间距 0.05 μm；Courant=0.25；关闭介电常数平均'],['计算区域','PML 厚 12 μm；上、下空气层各 12 μm；每例 16 个 MPI 进程'],['材料','Ordal 数据拟合的 Ti 稳定色散模型；未包含氧化层、加工粗糙度和温度依赖']],[1.45,5.55],9.5)
p('改变 F 时，表面开口由 25 增至 32.5、40 μm，最深处也由 40.054 增至 41.677、43.301 μm。因此这是一条由几何定义约束的参数路径，不是“仅改变孔径且深度完全不变”的实验。')
p('Aflux=1−R−T；Avol 为独立体积损耗积分得到的吸收率。将吸收率用于发射率解释时，需要对应方向和偏振的一致通道定义以及互易、局部热平衡假设。当前图表首先报告数值吸收响应，不代表实测热辐射功率。[R1]')

page('2 占空比呈现三点非单调趋势')
image('current/01_fill_comparison.png','图 2  10.5 μm、p 偏振、r20 的占空比对比。连接线只引导观察离散样点；Aflux 与 Avol 是两种计算量，不是误差棒。正负角差单位为百分点。[D1、D2、D3]',6.9)
p('F 从 0.50 增至 0.65 时，+30° Aflux 增加 6.646 个百分点；继续增至 0.80 时下降 3.870 个百分点。对应正负角差为 4.818、11.039、7.815 个百分点。体积积分给出的差值为 4.814、11.038、7.807 个百分点，三点排序一致。')
table(['F','+30° Avol','−30° Avol','Avol 正负角差'],[[f'{f:.2f}',pct(pair(f,10.5)[0]['metrics']['A_vol'])+'%',pct(pair(f,10.5)[1]['metrics']['A_vol'])+'%',f'{100*(pair(f,10.5)[0]["metrics"]["A_vol"]-pair(f,10.5)[1]["metrics"]["A_vol"]):.3f} pp'] for f in [.5,.65,.8]],[.8,1.9,1.9,2.4])
p('这一证据支持优先在 F≈0.65 附近细化，但峰位可能位于尚未采样的区域。当前数据不能区分腔模耦合、开口变化、有效深度或壁厚变化各自的贡献，不能把曲线起伏直接归因于某一种共振机制。')
p('−30° 下 F=0.50 与 0.80 的差仅为 0.222 个百分点，小于单点通量与体积吸收差的量级；这部分细微排序不宜作为可靠的设计结论。')

page('3 三个波长抽样均呈正负角不对称')
image('current/02_wavelength_samples.png','图 3  F=0.65 的 p 偏振离散波长抽样。三个波长分别计算，不构成连续 8–13 μm 光谱。[D2、D3]',6.9)
table(['波长','+30° Aflux','−30° Aflux','角差','比值'],[[f'{l:g} μm',pct(pair(.65,l)[0]['metrics']['A_flux'])+'%',pct(pair(.65,l)[1]['metrics']['A_flux'])+'%',f'{pair(.65,l)[2]:.3f} pp',f'{pair(.65,l)[0]["metrics"]["A_flux"]/pair(.65,l)[1]["metrics"]["A_flux"]:.3f}'] for l in [8,10.5,13]],[1.0,1.5,1.5,1.6,1.4])
p('三个样点中，+30° 响应均高于 −30°；随着所抽样波长增大，两侧吸收率均下降。比值约为 2.03–2.24，但它只是这两个角度的响应比，不是对整个半球积分定义的方向性因子。')
p('8 μm 与 13 μm 的对应平板精度检查仍存在超阈值问题，故不能只凭曲线顺滑或方向相同就把这些点提升为正式合格结果。10.5 μm 对应平板检查通过，也不自动证明所有结构和波长的精度。')
p('下一轮应补齐 8–13 μm 的谱角数据后再计算带内指标。没有指定工作温度和谱权重时，不应将三点平均称为热辐射加权发射率。固定横向波矢的一次宽带运行还会使不同波长对应不同角度，不能未经处理直接当作固定角光谱。[R2]')

page('4 数值偏差及其影响')
image('current/03_numerical_residuals.png','图 4  10 个当前结构样点的通量与体积吸收差，以及保留符号的透射率。小负值属于尚待解释的数值残差，不代表物理负透射。[D1–D3]',6.9,max_height=4.1)
p('全部 10 点的 |Aflux−Avol| 为 0.269–0.380 个百分点；T 均有小负值，绝对值最大约为入射功率的 0.08354%。这些偏差小于主要正负角差，但不是完整的不确定度估计，也不能据此保证峰位、微弱变化或绝对发射率准确。')
table(['问题','已知影响','处理原则'],[['吸收计算差','影响绝对 A；可能干扰小幅参数排序','同时保留 Aflux 与 Avol，不加统一常数修正'],['小负透射','阻碍严格被动性与能量解释','保留符号；检查总场、归一化与返回波'],['平板校准不全','8、13 μm 等点仍有校准偏差','保留探索标签，验证关键候选再做定量结论'],['结构 s 偏振缺失','不能报告非偏振或双偏振方向性','在一致源定义下补算同结构 s 通道']],[1.25,2.25,3.5],9.5)
p('目前三点占空比的主要起伏在两种吸收计算中一致，可用于筛选；若要比较相差仅几十分之一个百分点的候选，需要先降低或量化上述偏差。')

page('5 源类型诊断解释了部分偏振差异')
image('current/04_source_type_diagnostic.png','图 5  正入射、10.5 μm、r20 平板的源类型对照。比较的是电流源与磁流源实现，不是斜槽的实际 s/p 热发射差异。[D6]',6.9)
p('采用相同电流源类型时，平板 s/p 的 Aflux 差约 5.55×10⁻¹⁶；采用相同磁流源类型时约 3.33×10⁻¹⁶。电流源与磁流源之间的 Aflux 相差 0.37087 个百分点，而 Avol 的差只有约 2.49×10⁻⁸（无量纲吸收率）。')
p('这一对照说明，所测正入射平板的“原生 p/s 分裂”主要随源类型变化，并非平板本身具有对应的物理偏振差异。不能把这项控制结论直接推广为斜槽没有偏振效应。')
p('保存场的功率分解还表明，入射场与散射场的相干交叉项可以重现吸收记账差的大部分。它解释了偏差如何进入通量结果，但尚未唯一确定返回波的来源，也没有消除体积吸收与解析平板值之间的剩余差异。[D7]')
p('下一步把“定位源与参考场机制”和“结构参数筛选”分开推进。对已保留的场重新计算功率分解，优先做少量关键对照；不要将正常 p/s 物理差异和实现差异混在同一张最终性能图中。')

page('6 平板验证与材料模型的边界')
flat_rows=[]
for x in flat:
    flat_rows.append([x['polarization'],f'{x["wavelength_um"]:g}',f'{x["theta_deg"]:+g}',x['resolution'],f'{100*abs(x["metrics"]["R"]-x["analytic"]["R"]):.4f}', '通过' if 'PASSED' in x['status'] else '未通过'])
image('current/06_flat_validation.png','图 6  当前材料模型下的平板反射率误差与 0.25 个百分点精度线。p 通道使用 r20，s 通道使用 r28；重复的 s 通道为不同源宽度对照，并非不同几何。[D5]',6.9)
p('该批 12 项平板控制中，4 项通过当前 0.25 个百分点的反射率精度阈值，8 项未通过。通过点是 p 偏振 10.5、20 μm 的 ±30°；未通过点包括 p 偏振 3、8、13 μm 的 ±30°以及两个源宽度的 s 偏振 10.5 μm 对照。两条 s 对照不是两个不同结构。[D5]')
sub('稳定材料模型已进入当前计算')
p('当前 10 个结构点均使用 SHA256 前缀 283197fec119 的 Ti 色散模型，背景介电常数约 23.7943。早先宽带拟合模型在 Courant=0.25 下的发散已由 16 个小域短测试复现；无界面、无 PML 的控制也可发散，不能归因于停止阈值太严。换用稳定模型后，当前结构批次已完成，但这不消除其他误差。[D8、R3]')
p('当前模型频域拟合覆盖 2.5–25 μm，目标工作范围为 3–20 μm。拟合通过与时域收敛、平板误差和真实 TA2 样品吻合是不同层次的验证。当前材料数据未对激光加工表面、氧化层和具体温度校准。')

page('7 历史几何离散诊断')
image('historical/H01_flat_discretization_20260922.png','图 H1  旧模型的平板空间离散诊断。r20、r24、r32 曲线是独立离散方程预测，不是已完成的 FDTD；不能用作当前斜槽网格的合格证明。[H1]',6.9,max_height=6.2)
p('历史诊断的价值是记录哪些数值因素已被检查、当时出现何种差异，以及为什么后续改变验证路线。它不能替代在当前材料、源设置和几何下重新确认关键候选的误差。')
p('当前 r20 的三组 F 对比是同一网格层级上的参数筛选。较早 r24、r28 或 r32 的图即使趋势相近，也不能跨材料版本拼接成当前结果的收敛曲线。')

page('8 历史边界与尾场诊断')
image('historical/H02_boundary_screen_20260924.png','图 H2  既有窄域边界筛选图。三项原算例均仍为 NUMERICAL_FAILED；一项边界指标改善不能代表全宽斜槽或当前材料合格。[H2]',6.9,max_height=3.2)
image('historical/H03_r24_tail_reflectance_20260930.png','图 H3  既有 r24 保存场尾部与反射率诊断。r24 未达旧停止线；图中的 r16/r20 状态只属于历史单条件，不继承为当前材料的资格。[H3]',6.9,max_height=3.4)

page('9 下一步计算与目标距离')
p('目标已推进到“有可比较的参数筛选证据”，尚未达到“8–13 μm 双偏振带内方向性规律”。优先沿当前 F≈0.65 候选补齐必要维度，同时用少量对照控制误差，比扩大所有参数的组合更有针对性。')
table(['优先级','建议任务','应得到的证据'],[['1','核对关键样点的总场功率、参考场与源类型；关键点网格或时间步复核','偏差量级及原因的可重复估计，保留未解释残差'],['2','F=0.65，补 s 偏振 8、10.5、13 μm 的 ±30°，使用一致源协议','s/p 对比和非偏振合成的基础'],['3','在 8–13 μm 加密谱点，再对候选角域加密角度','带平均响应、带内正负角差、峰值方向和角宽'],['4','在已验证候选附近细化 F，并独立控制轴长或深度','分离几何耦合，定位局部最优区域'],['5','稳定完成主目标后扩展 3–20 μm；再进入三维及样品校准','带外响应、二维向实物的适用边界']],[.65,3.2,3.15],10)
sub('建议统一的后续指标')
p('同时报告每个角度的 8–13 μm 带内平均响应和完整谱角图。选择普通波长平均还是温度相关的黑体加权，应在任务表中明确；非偏振响应由相同条件的 s、p 通道平均得到。正负角差、响应比与全角域方向性因子应分别命名，避免混用。')
p('机器学习暂不宜承担最优设计结论：当前同模型仅有 10 个结构点，维度覆盖不足且存在系统偏差。待形成具有一致质量标记的参数集后，可用于代理模型与主动选点，预测误差仍需独立计算验证。')
p('本次归档不改变停止阈值、不重启任务，也不启动新一轮求解。后续排程仍应尊重共享服务器上的其他使用者。')

page('10 关键数据总表')
table(['F','λ μm','θ °','Aflux %','Avol %','T %','差 pp'],[[f'{x["fill"]:.2f}',f'{x["wavelength_um"]:g}',f'{x["theta_deg"]:+g}',pct(x['metrics']['A_flux']),pct(x['metrics']['A_vol']),f'{100*x["metrics"]["T"]:.5f}',f'{100*(x["metrics"]["A_flux"]-x["metrics"]["A_vol"]):.3f}'] for x in sorted(records,key=lambda x:(x['fill'],x['wavelength_um'],-x['theta_deg']))],[.55,.65,.65,1.35,1.35,1.25,1.2],9.5)
p('表中差值为 Aflux−Avol。所有数据均保留原始符号；百分数为无量纲结果乘 100，百分点为两个百分数之差。未通过正式资格认定的结果仍按探索数据保存。完整数值精度、检查项和来源哈希见数据包。')
sub('本轮探索性停止条件')
p('参考计算使用强度衰减 10⁻⁴；结构计算使用 10⁻³。探针以各自历史峰值归一化，连续 3 个窗口通过，同时要求 R/T 绝对漂移小于 0.001（0.1 个百分点）；源结束后至少再计算 300 个 Meep 时间单位。探索性吸收差和级次闭合阈值为 0.01（1 个百分点），参考伪反射阈值为 0.001（0.1% 入射功率）。[D4]')
p('上述阈值是当前筛选协议，不是普适的物理真理。“检查通过”意味着满足本轮明确的容差，不能据此声称误差小于任意更严格目标。Aflux 与 Avol 的差也不等价于全部数值误差。')

page('11 图表归档与来源')
p('数据包分为当前同模型结果、独立控制与历史图三部分。当前主图由保存的结果重新绘制；历史图保持原图并逐项标注。每幅图的来源、适用范围及哈希在图表目录中记录。')
table(['内容','包内位置'],[['当前结构及角度对','data/current_structured_results.json；data/direction_pairs.json'],['控制和材料','data/flat_validation_results.json；data/source_type_diagnostic.json；data/current_model.json'],['几何和计算协议','data/geometry_and_settings.json'],['当前重绘图','figures/current/（PNG 与 SVG）'],['既有图表与去重目录','figures/historical/；data/existing_figure_catalog.json'],['再绘图及完整性','scripts/plot_results.py；manifest.sha256']],[1.65,5.35],9.5)
p('D1 最新四项完成核查，2026-10-14；D2 同模型 F=0.65 基线核查；D3 8/13 μm 完成核查，2026-10-13；D4 几何与锁定参数；D5 14 项夜间验证中的 12 个平板控制；D6 正入射源类型对照；D7 保存场功率分解；D8 材料稳定性小域诊断。数据文件的 source_manifest 给出原证据文件标识及 SHA256。')
p('H1–H3 的完整来源与原图哈希见 existing_figure_catalog.json。其他代表性历史图同包归档，不作为当前结果的数值验证。未上传 SSH 凭据、主机与网络信息、调度脚本、完整原始场文件或环境磁盘镜像。')
sub('方法参考')
p('R1  J J Greffet 等，Light Emission by Nonequilibrium Bodies: Local Kirchhoff Law，Physical Review X 8, 021008 (2018)，https://doi.org/10.1103/PhysRevX.8.021008。方向及局部热平衡条件需与当前计算通道对应。', 'Caption')
p('R2  Meep FAQ，Gaussian source 与角度光谱，以及负反射或透射的可能原因。https://meep.readthedocs.io/en/latest/FAQ/ （2026-10-14 查阅）。', 'Caption')
p('R3  Meep Materials，Numerical Stability。https://meep.readthedocs.io/en/latest/Materials/#numerical-stability （2026-10-14 查阅）。', 'Caption')
p('材料来源标识：Ordal 等，Applied Optics 27, 1203 (1988)，DOI 10.1364/AO.27.001203；数据库转录及拟合来源信息保留在 current_model.json。模型尚未对具体加工样品验证。','Caption')

catalog=read('existing_figure_catalog.json')['figures']
for title,ids in [('附录一 早期吸收与尾场图',['H04','H05']),('附录二 边界灵敏度与旧材料图',['H06','H07']),('附录三 金属场与历史网格图',['H08','H09']),('附录四 三维几何与平板控制明细',['H10'])]:
    page(title)
    for ident in ids:
        item=next(x for x in catalog if x['id']==ident)
        image(str(Path(item['file']).relative_to('figures')),ident+'  '+item['caption_zh'],6.9,max_height=2.7 if ident=='H10' else 3.3)
    if ids==['H10']:
        table(['偏振','λ μm','θ °','r','R 偏差 pp','平板精度'],flat_rows,[.7,.8,.75,.6,1.7,2.05],8.5)

target=OUT/'钛斜槽定向热发射计算阶段报告_20261014.docx'
doc.save(target)
print(target)
